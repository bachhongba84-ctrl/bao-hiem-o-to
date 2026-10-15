import base64
import hashlib
import hmac
import io
import json
import re
import time
from concurrent.futures import ThreadPoolExecutor, as_completed

import requests
import pandas as pd
import streamlit as st
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from PIL import Image, ImageOps

st.set_page_config(page_title="Đọc ảnh bảo hiểm ô tô", page_icon="🚗", layout="wide")

FIELDS = [
    "Phí bảo hiểm mục 7",
    "Tổng phí bảo hiểm mục 8",
    "Chủ xe cơ giới",
    "Địa chỉ",
    "Điện thoại",
    "Số biển kiểm soát",
    "Số khung",
    "Số máy",
    "Năm sản xuất",
    "Loại xe",
    "Trọng tải",
    "Số chỗ ngồi",
    "Mục đích sử dụng",
    "Thời hạn bảo hiểm",
    "Phí bảo hiểm mục 4",
    "Ngày cấp",
    "Số giấy chứng nhận",
]
PURPOSE_OPTIONS = ["Chưa xác định", "Kinh doanh", "Không kinh doanh"]
COLS = ["Xóa", "Nguồn ảnh", "Cần kiểm tra"] + FIELDS
DEFAULT_MODEL = "gemini-2.5-flash"

PROMPT = (
    "Các ảnh trên là các trang của MỘT giấy chứng nhận bảo hiểm bắt buộc "
    "trách nhiệm dân sự của chủ xe cơ giới (ô tô) ở Việt Nam. "
    "Hãy đọc kỹ và trích xuất đúng các trường sau, chép nguyên văn như "
    "trên giấy (giữ nguyên số, dấu chấm, dấu phẩy):\n"
    + "\n".join(f"- {f}" for f in FIELDS)
    + "\n\nQuy tắc:\n"
    "- 'Phí bảo hiểm mục 4', 'Phí bảo hiểm mục 7', 'Tổng phí bảo hiểm mục 8' "
    "là số tiền ở các mục đánh số 4, 7, 8 trên giấy.\n"
    "- 'Số giấy chứng nhận' thường ở đầu trang 2.\n"
    "- 'Mục đích sử dụng' chỉ được là đúng một trong ba giá trị: "
    "'Kinh doanh', 'Không kinh doanh', 'Chưa xác định'. "
    "Chỉ chọn 'Kinh doanh' hoặc 'Không kinh doanh' khi thấy rõ ô nào được "
    "đánh dấu tích; nếu không chắc, ghi 'Chưa xác định'.\n"
    "- Không đoán. Trường nào không thấy hoặc không đọc rõ thì để chuỗi rỗng "
    "và ghi tên trường đó vào danh sách 'Can_kiem_tra'.\n"
    "- Trường nào đọc được nhưng chưa chắc chắn cũng ghi vào 'Can_kiem_tra'.\n\n"
    "Chỉ trả về MỘT đối tượng JSON, không thêm chữ nào khác, không dùng "
    "```. Các khóa đúng bằng tên các trường ở trên, cộng thêm khóa "
    "'Can_kiem_tra' là danh sách tên các trường cần kiểm tra."
)


def prep_image(data):
    """Làm thẳng chiều ảnh, thu nhỏ vừa đủ để gửi nhanh, trả về base64 JPEG."""
    img = Image.open(io.BytesIO(data))
    img = ImageOps.exif_transpose(img).convert("RGB")
    longest = max(img.size)
    if longest > 1800:
        scale = 1800 / longest
        img = img.resize(
            (max(1, int(img.width * scale)), max(1, int(img.height * scale))),
            Image.LANCZOS,
        )
    buf = io.BytesIO()
    img.save(buf, format="JPEG", quality=88)
    return base64.standard_b64encode(buf.getvalue()).decode("ascii")


def parse_json(text):
    m = re.search(r"\{.*\}", text, re.S)
    if not m:
        raise ValueError("AI không trả về dữ liệu đúng định dạng.")
    return json.loads(m.group(0))


def read_group(api_key, model, images_b64):
    parts = []
    for i, b64 in enumerate(images_b64, start=1):
        parts.append({"text": f"Ảnh {i}:"})
        parts.append({"inline_data": {"mime_type": "image/jpeg", "data": b64}})
    parts.append({"text": PROMPT})

    body = {
        "contents": [{"parts": parts}],
        "generationConfig": {"responseMimeType": "application/json", "temperature": 0},
    }
    url = f"https://generativelanguage.googleapis.com/v1beta/models/{model}:generateContent"
    headers = {"x-goog-api-key": api_key, "Content-Type": "application/json"}

    data = None
    last = ""
    for attempt in range(4):
        r = requests.post(url, headers=headers, json=body, timeout=120)
        if r.status_code in (429, 500, 503):
            last = f"HTTP {r.status_code}"
            time.sleep(6 * (attempt + 1))
            continue
        if r.status_code != 200:
            raise RuntimeError(f"HTTP {r.status_code}: {r.text[:300]}")
        data = r.json()
        break
    if data is None:
        raise RuntimeError(
            f"Google đang quá tải hoặc đã hết lượt miễn phí ({last}). "
            "Hãy thử lại sau vài phút."
        )

    cands = data.get("candidates") or []
    if not cands:
        raise ValueError("AI không trả kết quả (ảnh có thể không đọc được).")
    text = "".join(
        p.get("text", "") for p in cands[0].get("content", {}).get("parts", [])
    )
    parsed = parse_json(text)

    result = {}
    for f in FIELDS:
        v = parsed.get(f, "")
        result[f] = "" if v is None else str(v).strip()
    if result["Mục đích sử dụng"] not in PURPOSE_OPTIONS:
        result["Mục đích sử dụng"] = "Chưa xác định"
    check = parsed.get("Can_kiem_tra", [])
    result["Cần kiểm tra"] = (
        ", ".join(str(x) for x in check) if isinstance(check, list) else str(check)
    )
    return result


def excel_bytes(df_export):
    output = io.BytesIO()
    with pd.ExcelWriter(output, engine="openpyxl") as writer:
        df_export.to_excel(writer, index=False, sheet_name="Bao hiem")
        ws = writer.sheets["Bao hiem"]
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions
        for cell in ws[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill(fill_type="solid", fgColor="1F4E78")
            cell.alignment = Alignment(
                horizontal="center", vertical="center", wrap_text=True
            )
        ws.row_dimensions[1].height = 35
        for col in range(1, len(FIELDS) + 1):
            ws.column_dimensions[get_column_letter(col)].width = 24

        guide = writer.book.create_sheet("Huong dan")
        guide.append(["LƯU Ý KIỂM TRA"])
        guide.append(["Đối chiếu tất cả dữ liệu với ảnh gốc."])
        guide.append(["Kiểm tra số giấy, biển số, số khung, số máy."])
        guide.append(["Kiểm tra phí, thời hạn và ngày cấp."])
        guide.append(["Xác nhận dấu tích kinh doanh/không kinh doanh."])
        guide.column_dimensions["A"].width = 75
    return output.getvalue()


def login():
    st.title("🔐 Đăng nhập phần mềm")
    st.caption("Đọc ảnh giấy chứng nhận bảo hiểm ô tô")
    with st.form("login_form"):
        password = st.text_input("Mật khẩu", type="password")
        submitted = st.form_submit_button("Đăng nhập")
    if submitted:
        configured = st.secrets.get("PASSWORD_SHA256", "")
        actual = hashlib.sha256(password.encode("utf-8")).hexdigest()
        if configured and hmac.compare_digest(actual, configured.strip().lower()):
            st.session_state["authenticated"] = True
            st.rerun()
        else:
            st.error("Mật khẩu không đúng hoặc chưa cấu hình Secrets.")


def file_key(f):
    return f"{f.name}|{f.size}"


def reset_all():
    st.session_state["table"] = pd.DataFrame(columns=COLS)
    st.session_state["processed"] = set()
    st.session_state["grid_version"] = st.session_state.get("grid_version", 0) + 1


# ---------------- Khởi tạo ----------------
if "authenticated" not in st.session_state:
    st.session_state["authenticated"] = False
if not st.session_state["authenticated"]:
    login()
    st.stop()
if "table" not in st.session_state:
    reset_all()

api_key = st.secrets.get("GEMINI_API_KEY", "")
model = st.secrets.get("MODEL", DEFAULT_MODEL)

st.title("🚗 Đọc ảnh bảo hiểm ô tô")
st.write(
    "Chọn ảnh, bấm đọc, sửa trực tiếp trong bảng rồi tải Excel. "
    "AI đọc ảnh nên ảnh càng rõ càng chính xác."
)

with st.sidebar:
    st.subheader("Phiên làm việc")
    st.write(f"Đang có {len(st.session_state['table'])} dòng.")
    if st.button("Đăng xuất"):
        st.session_state["authenticated"] = False
        st.rerun()
    if st.button("Xóa toàn bộ danh sách"):
        reset_all()
        st.rerun()

if not api_key:
    st.error(
        "Chưa có khóa AI. Vào Streamlit > Settings > Secrets và thêm dòng "
        'GEMINI_API_KEY = "khóa của bạn".'
    )
    st.stop()

uploads = st.file_uploader(
    "Chọn ảnh giấy bảo hiểm (chọn được nhiều ảnh)",
    type=["jpg", "jpeg", "png", "webp"],
    accept_multiple_files=True,
)
per_group = st.selectbox(
    "Mỗi giấy bảo hiểm gồm mấy ảnh? (ảnh được ghép theo thứ tự chọn)",
    [1, 2, 3],
    index=1,
)

groups = []
if uploads:
    for i in range(0, len(uploads), per_group):
        groups.append(uploads[i : i + per_group])
    for n, g in enumerate(groups, start=1):
        st.caption(f"Giấy {n}: " + " + ".join(f.name for f in g))

new_groups = [
    g for g in groups if "||".join(file_key(f) for f in g) not in st.session_state["processed"]
]

if new_groups:
    if st.button(f"🔎 Đọc {len(new_groups)} giấy bằng AI", type="primary"):
        jobs = []
        for g in new_groups:
            key = "||".join(file_key(f) for f in g)
            names = " + ".join(f.name for f in g)
            jobs.append((key, names, [prep_image(f.getvalue()) for f in g]))

        bar = st.progress(0.0, text="Đang đọc...")
        new_rows, errors = [], []
        with ThreadPoolExecutor(max_workers=2) as pool:
            futures = {
                pool.submit(read_group, api_key, model, imgs): (key, names)
                for key, names, imgs in jobs
            }
            done = 0
            for fut in as_completed(futures):
                key, names = futures[fut]
                done += 1
                try:
                    res = fut.result()
                    st.session_state["processed"].add(key)
                except Exception as e:
                    res = {f: "" for f in FIELDS}
                    res["Mục đích sử dụng"] = "Chưa xác định"
                    res["Cần kiểm tra"] = "Không đọc được, hãy nhập tay"
                    errors.append(f"{names}: {e}")
                row = {"Xóa": False, "Nguồn ảnh": names}
                row.update(res)
                new_rows.append(row)
                bar.progress(done / len(futures), text=f"Đã đọc {done}/{len(futures)} giấy")
        bar.empty()

        add_df = pd.DataFrame(new_rows, columns=COLS)
        cur = st.session_state["table"]
        st.session_state["table"] = (
            pd.concat([cur, add_df], ignore_index=True) if len(cur) else add_df
        )
        for msg in errors:
            st.warning("Có lỗi: " + msg)
        st.rerun()

# ---------------- Bảng chỉnh sửa ----------------
table = st.session_state["table"]
if len(table):
    st.subheader("Kiểm tra và sửa thông tin")
    st.caption(
        "Bấm hai lần vào ô để sửa. Cột 'Cần kiểm tra' liệt kê các ô AI chưa chắc. "
        "Tích cột 'Xóa' rồi bấm nút bên dưới để bỏ dòng thừa."
    )
    edited = st.data_editor(
        table,
        key=f"grid_{st.session_state['grid_version']}",
        num_rows="fixed",
        hide_index=True,
        use_container_width=True,
        disabled=["Nguồn ảnh", "Cần kiểm tra"],
        column_config={
            "Xóa": st.column_config.CheckboxColumn("Xóa", width="small"),
            "Mục đích sử dụng": st.column_config.SelectboxColumn(
                "Mục đích sử dụng", options=PURPOSE_OPTIONS, required=True
            ),
        },
    )
    st.session_state["table"] = edited

    c1, c2 = st.columns([1, 3])
    with c1:
        if st.button("🗑️ Xóa các dòng đã tích"):
            keep = edited[~edited["Xóa"].astype(bool)].reset_index(drop=True)
            st.session_state["table"] = keep
            st.session_state["grid_version"] += 1
            st.rerun()
    with c2:
        st.download_button(
            "📥 Tải file Excel",
            data=excel_bytes(edited[FIELDS].fillna("")),
            file_name="Danh_sach_bao_hiem.xlsx",
            mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )
else:
    st.info("Chọn ảnh ở phía trên rồi bấm nút đọc, kết quả sẽ hiện thành bảng ở đây.")

if uploads:
    with st.expander("🔍 Xem ảnh gốc để đối chiếu"):
        chosen = st.selectbox("Chọn ảnh", [f.name for f in uploads])
        picked = next(f for f in uploads if f.name == chosen)
        st.image(picked, caption=picked.name)

st.caption("Dữ liệu do AI đọc có thể sai. Hãy đối chiếu ảnh gốc trước khi dùng.")
